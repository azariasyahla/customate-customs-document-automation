import base64, io, re
from copy import deepcopy
from datetime import datetime, timezone, timedelta
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image


RULES=[
(['HOSE, FUEL'],'rubber dan synthetic polymer','menyalurkan bahan bakar antar komponen','menjadi saluran fleksibel untuk aliran bahan bakar','menjaga kelancaran suplai bahan bakar','getaran dan perubahan bentuk'),
(['HOSE AIR CLEANER'],'rubber dan synthetic polymer','menyalurkan udara pada rangkaian komponen pemasukan udara','menghubungkan saluran pemasukan udara dengan komponen terkait','menjaga kelancaran aliran udara','getaran dan perubahan bentuk'),
(['HOSE, HIGH'],'rubber dan synthetic polymer','menyalurkan fluida bertekanan antar komponen','menghubungkan komponen melalui saluran fleksibel','menjaga kelancaran aliran fluida','tekanan dan perubahan bentuk'),
(['HOSE'],'rubber dan synthetic polymer','menyalurkan fluida atau udara antar komponen','menjadi saluran fleksibel antar komponen','menjaga kelancaran aliran','getaran dan perubahan bentuk'),
(['SEAL','O-RING','RING O','PACKING'],'rubber dan synthetic polymer','memberikan penyegelan pada sambungan komponen','membantu mencegah kebocoran dan masuknya kotoran','menjaga kekedapan sistem','tekanan dan perubahan bentuk'),
(['BELT ASSY SEAT'],'polymer berkekuatan tinggi','menahan posisi operator selama penggunaan','membantu menjaga operator tetap pada posisi duduk','mendukung keselamatan operator','tarikan dan penggunaan berulang'),
(['BELT V'],'rubber dan synthetic polymer','meneruskan tenaga putar antar komponen','menghubungkan putaran melalui sistem pulley','menjaga kelancaran transfer tenaga','gesekan dan penggunaan berulang'),
(['HORN'],'metal alloy','menghasilkan sinyal suara peringatan','memberikan tanda suara kepada lingkungan sekitar','mendukung keselamatan operasional','getaran dan penggunaan berulang'),
(['RADIATOR ASSY'],'metal alloy','melepaskan panas dari cairan pendingin','membantu mendinginkan cairan sebelum kembali ke sistem','menjaga temperatur kerja','perubahan temperatur dan penggunaan berulang'),
(['HOSE','SEAL','O-RING','RING O','PACKING'],'rubber dan synthetic polymer','menyalurkan atau menjaga kekedapan aliran pada rangkaian komponen','mendukung aliran atau penyegelan sesuai penggunaannya','menjaga kestabilan kerja sistem','tekanan dan perubahan bentuk'),
(['GASKET'],'gasket sheet berbasis compressed fiber','memberikan penyegelan pada sambungan komponen','mencegah kebocoran pada sambungan','menjaga kekedapan sistem','tekanan dan getaran'),
(['BEARING'],'metal alloy','menopang pergerakan komponen berputar','mengurangi gesekan dan menopang beban','menjaga kestabilan putaran','gesekan dan beban kerja'),
(['SHAFT','GEAR','VALVE','PIPE','LEVER','BRAKE SHOE','SHOE','CYLINDER','PUMP'],'metal alloy','mendukung kerja mekanis antar komponen','meneruskan gerakan atau gaya antar komponen','menjaga kelancaran operasional','beban kerja dan penggunaan berulang'),
(['SWITCH','SENSOR','KEYBOARD','DISPLAY','HARNESS','CABLE','BUTTON'],'engineering plastic','mendukung fungsi pengendalian atau penghubung pada rangkaian komponen','mendukung penyampaian kendali sesuai penggunaannya','menjaga kelancaran pengoperasian','penggunaan berulang'),
]

MATERIALS=['metal alloy','engineering plastic','polymer berkekuatan tinggi','rubber dan synthetic polymer','gasket sheet berbasis compressed fiber','carbon','glass']

EXTRA_RULES=[
(['CARBON'], 'carbon', 'mendukung fungsi komponen antar komponen', 'membantu kerja komponen lainnya', 'menjaga kelancaran kerja sistem', 'penggunaan berulang'),
(['GLASS','LENS'], 'glass', 'mendukung pengamatan atau penerusan cahaya', 'memberikan bidang transparan pada komponen terkait', 'mendukung pengamatan selama operasional', 'penggunaan normal'),
(['RING','ROLLER','DISC','COVER','MOTOR','TANK','COIL','RADIATOR'], 'metal alloy', 'menopang kerja komponen lainnya', 'menjalankan fungsi pendukung sesuai posisi pemasangannya', 'menjaga kestabilan operasional', 'beban kerja dan penggunaan berulang'),
(['PLASTIC','KNOB'], 'engineering plastic', 'mendukung fungsi pengoperasian komponen', 'membantu pengoperasian antar komponen', 'menjaga kelancaran penggunaan', 'penggunaan berulang'),
]

SPECIFIC_RULES=[
 (['SHAFT'],'metal alloy','meneruskan putaran','menghubungkan gerak putar antar komponen','menjaga penyaluran tenaga','beban putar dan penggunaan berulang'),
 (['GEAR'],'metal alloy','meneruskan putaran dan tenaga','menghubungkan gerakan melalui gigi yang saling bersinggungan','menjaga perpindahan tenaga','gesekan dan beban kerja'),
 (['VALVE'],'metal alloy','mengatur aliran fluida atau gas','membuka dan menutup jalur aliran','menjaga aliran tetap terkendali','tekanan dan penggunaan berulang'),
 (['PIPE'],'metal alloy','menyalurkan fluida atau gas','menghubungkan jalur aliran antar komponen','menjaga kelancaran aliran','tekanan dan getaran'),
 (['LEVER'],'metal alloy','menggerakkan mekanisme kendali','meneruskan gerakan tangan ke komponen yang dikendalikan','memudahkan pengoperasian','beban dan penggunaan berulang'),
 (['BRAKE SHOE','SHOE'],'metal alloy','membantu pengereman','menghasilkan gesekan untuk memperlambat atau menghentikan gerakan','menjaga pengereman tetap terkendali','gesekan dan penggunaan berulang'),
 (['CYLINDER'],'metal alloy','menghasilkan gerakan mekanis','mengubah tekanan fluida menjadi dorongan atau tarikan','menjaga pergerakan tetap stabil','tekanan dan beban kerja'),
 (['PUMP'],'metal alloy','mengalirkan fluida','memindahkan fluida menuju komponen yang membutuhkan','menjaga pasokan fluida','tekanan dan penggunaan berulang'),
 (['ROLLER'],'metal alloy','memandu pergerakan komponen','menopang gerakan bergulir dan mengurangi gesekan','menjaga gerakan tetap lancar','beban dan gesekan'),
 (['COIL'],'metal alloy','menghasilkan medan elektromagnetik','mengaktifkan gerakan pada mekanisme solenoid','menjaga respons kendali','penggunaan berulang'),
]

def draft(name,index=0):
 mat='metal alloy';main='menopang kerja komponen lainnya';op='membantu kerja komponen lainnya';benefit='menjaga kelancaran operasional';res='penggunaan berulang'
 for words,m,f,o,v,r in SPECIFIC_RULES+RULES+EXTRA_RULES:
  if any(w in name.upper() for w in words):mat,main,op,benefit,res=m,f,o,v,r;break
 label=name.lower()
 openings=[
  f'{label.capitalize()} berbahan {mat} digunakan untuk {main}.',
  f'Material {mat} digunakan pada {label} untuk {main}.',
  f'Terbuat dari {mat}, {label} digunakan untuk {main}.',
  f'{label.capitalize()} menggunakan material {mat} untuk {main}.',
  f'Berbahan dasar {mat}, {label} berperan untuk {main}.',
  f'Bahan utama {label} adalah {mat}.',
  f'{label.capitalize()} dibuat dari material {mat} untuk {main}.',
  f'Sebagai komponen untuk {main}, {label} menggunakan bahan {mat}.',
  f'Pada {label}, material {mat} digunakan sebagai bahan utama.',
  f'{label.capitalize()} merupakan komponen berbahan {mat} untuk {main}.',
  f'Komponen {label} terbuat dari {mat}.',
  f'Material {mat} menjadi bahan utama {label}.',
  f'Untuk {main}, {label} dibuat menggunakan material {mat}.',
  f'{label.capitalize()} memanfaatkan material {mat} untuk {main}.',
  f'Komponen {label} berbahan {mat} digunakan untuk {main}.',
  f'{label.capitalize()} dibuat menggunakan bahan {mat}.',
  f'{label.capitalize()} tersusun dari material {mat} dan digunakan untuk {main}.',
  f'Material pembentuk {label} adalah {mat}.',
  f'{label.capitalize()} memakai bahan {mat} untuk {main}.',
  f'Dibuat menggunakan {mat}, {label} berperan untuk {main}.',
 ]
 middles=[
  f'Komponen ini berfungsi untuk {op} sehingga membantu {benefit}.',
  f'Perannya adalah {op} agar dapat {benefit}.',
  f'Komponen ini digunakan untuk {main}. Fungsinya adalah {op} sehingga membantu {benefit}.',
  f'Fungsi komponen ini adalah {op}. Peran tersebut membantu {benefit}.',
  f'Komponen ini membantu {op} sekaligus {benefit}.',
 ]
 endings=[
  f'Ketahanan terhadap {res} membantu komponen tetap bekerja dengan baik selama digunakan.',
  f'Ketahanannya terhadap {res} mendukung keandalan komponen selama pemakaian.',
  f'Material tersebut mendukung ketahanan komponen terhadap {res} dalam pemakaian sehari-hari.',
  f'Ketahanan bahan terhadap {res} membantu mempertahankan fungsi komponen.',
  f'Ketahanannya terhadap {res} turut mendukung kelancaran pengoperasian.',
 ]
 return mat,openings[index%20]+' '+middles[index%5]+' '+endings[(index//4)%5]

def add_formatted_function(paragraph,text,part_name,formatting):
 terms=['metal alloy','engineering plastic','rubber','synthetic polymer','gasket sheet','compressed fiber','carbon','glass','hose','o-ring','seal','bearing','shaft','gear','valve','pipe','lever','brake shoe','cylinder','pump','switch','sensor','keyboard','display','harness','cable','button','roller','belt','radiator','horn','motor','coil','ring','gasket','packing','coolant','intake','overhaul','fan','mast','king pin']
 terms.append(part_name.strip())
 pattern=re.compile(r'(?<![A-Za-z])(?:'+'|'.join(re.escape(t) for t in sorted(set(terms),key=len,reverse=True) if t)+r')(?![A-Za-z])',re.I)
 cursor=0
 def add(value,italic):
  if not value:return
  run=paragraph.add_run(value)
  if formatting is not None:run._r.insert(0,deepcopy(formatting))
  run.bold=False;run.italic=italic
 for match in pattern.finditer(text):
  add(text[cursor:match.start()],False);add(match.group(),True);cursor=match.end()
 add(text[cursor:],False)

def parse_doc(data):
 d=Document(io.BytesIO(data)); parts=[]
 for ti,t in enumerate(d.tables):
  header=None
  for ri,row in enumerate(t.rows):
   vals=[c.text.strip() for c in row.cells]
   upper=[v.upper() for v in vals]
   if 'PART NUMBER' in upper and 'FUNCTION' in upper:
    header={key:upper.index(key) for key in ['PART NUMBER','DESCRIPTION','FUNCTION']};continue
   if header is None or len({id(c._tc) for c in row.cells})<3: continue
   pn=vals[header['PART NUMBER']]; name=vals[header['DESCRIPTION']]
   if not pn or pn=='Detail Picture':continue
   image=''; picrow=None
   if ri+1<len(t.rows) and 'Detail Picture' in t.rows[ri+1].cells[0].text:
    picrow=ri+1
    blips=t.rows[picrow]._tr.xpath('.//a:blip')
    if blips:
     rid=blips[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
     img=d.part.related_parts[rid]
     image='data:'+img.content_type+';base64,'+base64.b64encode(img.blob).decode()
   mat,txt=draft(name,len(parts))
   parts.append(dict(id=len(parts),pn=pn,name=name,function=vals[header['FUNCTION']] or txt,material=mat,image=image,reviewed=False,table=ti,row=ri,col=header['FUNCTION'],picrow=picrow,source=''))
 if not parts:raise ValueError('A table containing Part Number, Description, and Function was not found.')
 return d,parts

def jakarta_document_date():
 now=datetime.now(timezone(timedelta(hours=7)))
 months=['Januari','Februari','Maret','April','Mei','Juni','Juli','Agustus','September','Oktober','November','Desember']
 return f'Jakarta, {now.day} {months[now.month-1]} {now.year}'

def update_letter_layout(doc):
 # Keep section margins supplied by the template. Align the signature inside them.
 body=next((p for p in doc.paragraphs if p.text.strip().lower()=='dengan hormat,'),None)
 signature=next((i for i,p in enumerate(doc.paragraphs) if p.text.strip().lower()=='hormat kami'),None)
 if signature is not None:
  section=doc.sections[0];usable=section.page_width-section.left_margin-section.right_margin
  block_width=min(usable,Cm(8.5))
  for p in doc.paragraphs[signature:]:
   text=p.text.strip()
   if text.startswith('DETAIL GAMBAR'):break
   if not text:continue
   if not (text.lower()=='hormat kami' or text.startswith('(') or 'dept' in text.lower() or 'department' in text.lower()):continue
   if body is not None:p.style=body.style
   p.paragraph_format.left_indent=Cm(0)
   p.paragraph_format.right_indent=Cm(0)
   p.paragraph_format.first_line_indent=Cm(0)
   p.alignment=WD_ALIGN_PARAGRAPH.LEFT
   p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(3)
   p.paragraph_format.line_spacing=1
   for run in p.runs:
    run.font.size=Pt(11)
    if body is not None and body.runs:
     if body.runs[0].font.name:run.font.name=body.runs[0].font.name
   # Strip leading spaces without replacing run formatting.
   for run in p.runs:
    if run.text.strip():run.text=run.text.lstrip();break
    run.text=''
 for p in doc.paragraphs:
  if re.match(r'^\s*Jakarta\s*,\s*\d',p.text,re.I):
   fmt=next((deepcopy(r._r.rPr) for r in p.runs if r._r.rPr is not None),None)
   for run in list(p.runs):p._p.remove(run._r)
   r=p.add_run(jakarta_document_date())
   if fmt is not None:r._r.insert(0,fmt)
   r.font.size=Pt(11)
   break

def export_doc(d,parts,changes):
 if len(changes)!=len(parts):raise ValueError('The item list is incomplete.')
 for original,change in zip(parts,changes):
  if change.get('id')!=original['id']:raise ValueError('The item order has changed.')
  if not change.get('function','').strip():raise ValueError('Function is empty for '+original['pn'])
 # Clone to preserve the uploaded document across retries.
 buf=io.BytesIO();d.save(buf); out=Document(io.BytesIO(buf.getvalue()))
 update_letter_layout(out)
 for original,c in zip(parts,changes):
  table=out.tables[original['table']]; cell=table.rows[original['row']].cells[original['col']]
  p=cell.paragraphs[0]; fmt=None
  for existing_run in p.runs:
   if existing_run._r.rPr is not None:
    fmt=deepcopy(existing_run._r.rPr);break
  if fmt is None:
   desc=table.rows[original['row']].cells[2]
   for paragraph in desc.paragraphs:
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
     fmt=deepcopy(paragraph.runs[0]._r.rPr);break
  for child in list(p._p):
   if child.tag.endswith('}r'):p._p.remove(child)
  add_formatted_function(p,c['function'].strip(),original['name'],fmt)
  for extra in cell.paragraphs[1:]:extra._p.getparent().remove(extra._p)
  if not c.get('image'):continue
  if original['picrow'] is None:raise ValueError('Detail Picture row not found for '+original['pn'])
  pc=table.rows[original['picrow']].cells[0]
  # Retain the existing label paragraph and its original runs/styles.
  label_paragraph=next((paragraph for paragraph in pc.paragraphs if 'detail picture' in paragraph.text.lower()),None)
  if label_paragraph is None:label_paragraph=pc.paragraphs[0]
  for paragraph in list(pc.paragraphs):
   if paragraph._p is not label_paragraph._p:paragraph._p.getparent().remove(paragraph._p)
  for drawing in list(label_paragraph._p.xpath('.//w:drawing')):
   drawing.getparent().remove(drawing)
  if not label_paragraph.text.strip():
   label_paragraph.add_run('Detail Picture')
  raw=base64.b64decode(c['image'].split(',',1)[1],validate=True)
  im=Image.open(io.BytesIO(raw)); im.verify(); im=Image.open(io.BytesIO(raw)).convert('RGB');im.thumbnail((1600,1200))
  b=io.BytesIO();im.save(b,format='PNG'); b.seek(0)
  width,height=im.size; scale=min(5.4/width,2.7/height)
  picture_paragraph=pc.add_paragraph()
  picture_paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
  picture_paragraph.add_run().add_picture(b,width=Inches(width*scale),height=Inches(height*scale))
  spacer=pc.add_paragraph()
  spacer.paragraph_format.space_before=Pt(0)
  spacer.paragraph_format.space_after=Pt(0)
  spacer.paragraph_format.line_spacing=Pt(12)
  spacer.add_run().font.size=Pt(12)
 b=io.BytesIO();out.save(b);return b.getvalue()
