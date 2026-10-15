import base64,io,tempfile,unittest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from demo import create_demo
from customate_core import parse_doc,export_doc,draft,jakarta_document_date

class DocumentTests(unittest.TestCase):
 def test_roundtrip(self):
  with tempfile.TemporaryDirectory() as folder:
   template,image=create_demo(folder);doc,items=parse_doc(template.read_bytes())
   self.assertEqual([p['pn'] for p in items],['DEMO-001','DEMO-002','DEMO-003'])
   for item in items:item['image']='data:image/png;base64,'+base64.b64encode(image.read_bytes()).decode()
   out=Document(io.BytesIO(export_doc(doc,items,items)))
   self.assertEqual(len(out.inline_shapes),3)
   self.assertIn(jakarta_document_date(),[p.text for p in out.paragraphs])
   for ri in (2,4,6):
    ps=out.tables[0].rows[ri].cells[0].paragraphs
    self.assertEqual(ps[0].runs[0].font.name,'Calibri')
    self.assertEqual(ps[-2].alignment,WD_ALIGN_PARAGRAPH.CENTER)
    self.assertEqual(ps[-1].text,'')
 def test_fonts_and_italics(self):
  for font in ('Calibri','Times New Roman'):
   with self.subTest(font=font),tempfile.TemporaryDirectory() as folder:
    template,_=create_demo(folder);doc=Document(template)
    for ri in (1,3,5):
     r=doc.tables[0].rows[ri].cells[3].paragraphs[0].add_run('');r.font.name=font;r.font.size=Pt(11)
    b=io.BytesIO();doc.save(b);doc,items=parse_doc(b.getvalue())
    out=Document(io.BytesIO(export_doc(doc,items,items)));runs=out.tables[0].rows[1].cells[3].paragraphs[0].runs
    self.assertTrue(all(r.font.name==font and r.font.size.pt==11 for r in runs))
    self.assertTrue(any(r.italic and 'rubber' in r.text for r in runs))
 def test_variants(self):
  texts=[]
  for i in range(20):
   mat,text=draft('PIPE SUB-ASSY',i);texts.append(text)
   self.assertIn(mat,text.split('.')[0]);self.assertNotIn('sistem terkait',text);self.assertNotIn('. Dengan',text)
  self.assertEqual(len(set(texts)),20)
 def test_preserve_existing_text(self):
  with tempfile.TemporaryDirectory() as folder:
   template,_=create_demo(folder);doc=Document(template);doc.tables[0].rows[1].cells[3].text='Reviewed text.'
   b=io.BytesIO();doc.save(b);_,items=parse_doc(b.getvalue());self.assertEqual(items[0]['function'],'Reviewed text.')
if __name__=='__main__':unittest.main()
