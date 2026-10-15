from pathlib import Path
from docx import Document
from docx.shared import Pt
from PIL import Image, ImageDraw

def create_demo(folder='demo_output'):
    folder=Path(folder);folder.mkdir(parents=True,exist_ok=True)
    doc=Document();doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(11)
    for text in ['SURAT PERNYATAAN','Jakarta, 1 Januari 2020','Kepada Yth. Petugas Pemeriksaan Barang','Contoh dokumentasi barang untuk demonstrasi aplikasi.','Terima kasih.','Hormat Kami','','(Nama Penandatangan)','Parts Documentation Dept.']:
        doc.add_paragraph(text)
    doc.add_page_break();doc.add_paragraph('DETAIL GAMBAR DAN SPESIFIKASI BARANG')
    table=doc.add_table(rows=1,cols=4);table.style='Table Grid'
    for cell,text in zip(table.rows[0].cells,['Item No','Part Number','Description','Function']):cell.text=text
    records=[('DEMO-001','HOSE'),('DEMO-002','BEARING'),('DEMO-003','PIPE SUB-ASSY')]
    for i,(pn,name) in enumerate(records,1):
        for cell,text in zip(table.add_row().cells,[str(i),pn,name,'']):cell.text=text
        cell=table.add_row().cells[0];cell.merge(table.rows[-1].cells[3]);cell.text='Detail Picture'
        for r in cell.paragraphs[0].runs:r.font.name='Calibri';r.font.size=Pt(11)
    template=folder/'demo_template.docx';doc.save(template)
    im=Image.new('RGB',(640,400),'#eaf0f8');draw=ImageDraw.Draw(im)
    draw.text((35,160),'DEMO IMAGE — REPLACE WITH A VERIFIED PART PHOTO',fill='#102b52')
    image=folder/'demo_placeholder.png';im.save(image)
    return template,image

if __name__=='__main__':
    print(create_demo())
